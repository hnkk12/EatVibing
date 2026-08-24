# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14 if level==1 else (10 if level==2 else 6))
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(26, 46, 26) # Dark Forest Green
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(40, 80, 40)
    elif level == 3:
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(60, 60, 60)
    return h

def add_callout(doc, title, text, bg_hex="F4F6F0", border_hex="4D5D1C"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(40, 70, 30)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(50, 50, 50)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def create_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "2D3748") # Dark Slate
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=90, bottom=90, left=130, right=130)
            p = row_cells[col_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(51, 51, 51)
                
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def build_document():
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(34, 34, 34)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(5)

    # ----------------------------------------------------
    # BÌA / COVER PAGE
    # ----------------------------------------------------
    p_org = doc.add_paragraph()
    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_org = p_org.add_run("BÁO CÁO THUYẾT MINH ĐỀ TÀI NGHIÊN CỨU KHOA HỌC & DỰ ÁN CÔNG NGHỆ (AISC 2026)\n")
    r_org.bold = True
    r_org.font.size = Pt(11)
    r_org.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph().paragraph_format.space_before = Pt(25)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("EATVIBING\n")
    r_title.bold = True
    r_title.font.size = Pt(28)
    r_title.font.color.rgb = RGBColor(45, 80, 22)
    
    r_subtitle = p_title.add_run("HỆ THỐNG TRỢ LÝ AI ĂN UỐNG CÁ NHÂN HÓA VÀ HỖ TRỢ RA QUYẾT ĐỊNH ẨM THỰC TỰ ĐỘNG THEO VÒNG LẶP PANTRY-TO-PLATE\n")
    r_subtitle.bold = True
    r_subtitle.font.size = Pt(14)
    r_subtitle.font.color.rgb = RGBColor(30, 41, 59)
    
    r_en = p_title.add_run("(EatVibing – Personalized AI Food Coach & Culinary Decision System)\n")
    r_en.italic = True
    r_en.font.size = Pt(11)
    r_en.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph().paragraph_format.space_before = Pt(30)
    
    add_callout(doc, "ĐỊNH VỊ SẢN PHẨM CỐT LÕI (ONE-SENTENCE PITCH)", 
        "EatVibing là trợ lý AI ăn uống cá nhân hóa, học từ nguyên liệu sẵn có, hồ sơ, khẩu vị và hành vi nấu thực tế để tạo hướng dẫn bữa ăn đáng tin cậy theo vòng lặp Pantry → Plan → Trust → Cook → Learn.\n\n"
        "Hệ thống kết hợp sức mạnh của Mô hình Ngôn ngữ Lớn (EatVibing AI Coach) hỗ trợ giao tiếp Text & Voice, cùng Lớp ủy thác ẩm thực (Culinary Trust Layer với trạng thái SAFE / REVIEW / BLOCK) mang tính xác định, giúp loại bỏ 100% rủi ro dị ứng và ảo giác dữ liệu dinh dưỡng.",
        bg_hex="F0FDF4", border_hex="16A34A")
        
    doc.add_page_break()

    # ----------------------------------------------------
    # CHƯƠNG 1: BỐI CẢNH VÀ TÍNH CẤP THIẾT
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 1: BỐI CẢNH, VẤN ĐỀ VÀ TÍNH CẤP THIẾT CỦA ĐỀ TÀI", level=1)
    
    add_styled_heading(doc, "1.1. Bối cảnh thực tiễn và Điểm nghẽn thị trường", level=2)
    doc.add_paragraph(
        "Trong nhịp sống hiện đại, việc duy trì một chế độ ăn uống lành mạnh, tiết kiệm và phù hợp với thể trạng cá nhân đang đối mặt với các rào cản nghiêm trọng:"
    )
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Mệt mỏi vì phải ra quyết định ('Hôm nay ăn gì?'): ").bold = True
    p.add_run("Người dùng mất từ 15 - 30 phút mỗi ngày chỉ để suy nghĩ thực đơn, gây ra tình trạng decision fatigue và xu hướng thỏa hiệp với đồ ăn nhanh chế biến sẵn nghèo nàn dinh dưỡng.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Lãng phí thực phẩm tại hộ gia đình (Household Food Waste): ").bold = True
    p.add_run("Khoảng 30-40% thực phẩm trong tủ lạnh bị vứt bỏ vì người dùng quên hạn sử dụng hoặc không biết cách kết hợp nguyên liệu lẻ tẻ còn sót lại thành một món ăn ngon.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Sự đứt gãy giữa Công thức - Dinh dưỡng - Thực tế nấu: ").bold = True
    p.add_run("Các ứng dụng hiện nay hoặc chỉ là kho công thức tĩnh (không biết trong tủ lạnh có gì), hoặc là ứng dụng đếm calo khô khan, hoặc chatbot AI tự do dễ bị 'ảo giác' (hallucination) tự bịa số liệu dinh dưỡng và nguyên liệu không an toàn.")

    add_styled_heading(doc, "1.2. Mục tiêu nghiên cứu và Chỉ số đo lường thành công", level=2)
    doc.add_paragraph(
        "Đề tài tập trung xây dựng EatVibing AI Coach đồng hành liên tục, kết hợp Lớp ủy thác Culinary Trust Layer (SAFE/REVIEW/BLOCK) và vòng lặp Pantry-to-Plate tự động."
    )
    
    metrics_headers = ["Nhóm chỉ số", "Chỉ số cụ thể", "Mục tiêu kỳ vọng"]
    metrics_data = [
        ["Tương tác & Chuyển đổi", "Guideline Acceptance Rate", "> 65% gợi ý trong ngày được người dùng chấp thuận"],
        ["Tương tác & Chuyển đổi", "Pick-to-Cook Conversion", "> 50% món được chọn chuyển thành hành động nấu thực tế"],
        ["Tương tác & Chuyển đổi", "Repeat-Cook Rate", "Tăng trưởng 30% tỷ lệ nấu lại các món hợp khẩu vị"],
        ["Tương tác & Chuyển đổi", "Voice Task Completion", "> 85% tác vụ rảnh tay hoàn thành chính xác bằng giọng nói"],
        ["Độ tin cậy & An toàn", "Allergen & Safety Pass Rate", "100% không để lọt món chứa chất gây dị ứng (Strict BLOCK)"],
        ["Độ tin cậy & An toàn", "Deterministic Nutrition Accuracy", "100% số liệu calo/macro được tính từ cơ sở dữ liệu xác thực"],
        ["Cá nhân hóa & Xã hội", "Food Waste Reduction", "Giảm 35% lượng thực phẩm bị bỏ quên quá hạn trong tủ lạnh"],
        ["Cá nhân hóa & Xã hội", "Cold-start to Personalized Quality", "Đạt độ khớp khẩu vị > 80% chỉ sau 3 lần nấu đầu tiên"]
    ]
    create_table(doc, metrics_headers, metrics_data, col_widths=[1.8, 2.5, 2.2])

    # ----------------------------------------------------
    # CHƯƠNG 2: TÍNH MỚI VÀ SO SÁNH GIẢI PHÁP
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 2: TÍNH MỚI, ĐIỂM ĐỘT PHÁ VÀ SO SÁNH GIẢI PHÁP", level=1)
    
    add_styled_heading(doc, "2.1. Tính mới và Điểm đột phá (Novelty)", level=2)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Persistent Personal Food Model: ").bold = True
    p.add_run("Lưu trữ liên tục trạng thái nguyên liệu (Smart Pantry), lịch sử nấu, khẩu vị chi tiết và phản hồi sau bữa ăn vào một hồ sơ học tập bền vững qua nhiều phiên.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Kiến trúc Tách biệt Trách nhiệm (LLM vs Deterministic Layer): ").bold = True
    p.add_run("LLM chỉ xử lý hội thoại tự nhiên và sáng tạo gợi ý; toàn bộ tính toán dinh dưỡng, quy đổi đơn vị, trừ kho và lọc dị ứng do Deterministic Policy Engine xử lý.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Qualified Engagement Ranking cho Trending: ").bold = True
    p.add_run("Xếp hạng món thịnh hành dựa trên hành vi nấu thực tế, đánh giá và nấu lại. Quan trọng nhất: Trending KHÔNG BAO GIỜ được vượt qua các ràng buộc dị ứng cá nhân!")

    add_styled_heading(doc, "2.2. So sánh với các giải pháp hiện hữu (Gap Analysis)", level=2)
    gap_headers = ["Tiêu chí so sánh", "Recipe App (Cookpad, Tasty)", "Calorie App (MyFitnessPal)", "Chatbot chung (ChatGPT)", "EatVibing (AI Food Coach)"]
    gap_data = [
        ["Nhận diện tủ lạnh", "Thủ công từ khóa", "Không hỗ trợ", "Nhập text tự do", "Smart Pantry (Text/Voice/Auto-sub)"],
        ["Cá nhân hóa thực đơn", "Tĩnh, danh mục cố định", "Thực đơn mẫu khô khan", "Sinh ngẫu nhiên mỗi lần", "Personalized Guideline thích ứng"],
        ["Độ chính xác Calo/Macro", "Thường không có", "Chuẩn nhưng nhập khó", "Dễ bịa số liệu (Ảo giác)", "Deterministic Calculation + USDA Grounding"],
        ["Kiểm soát An toàn & Dị ứng", "Tự đọc cảnh báo", "Cơ bản", "Không có lớp chặn cứng", "Culinary Trust Layer (SAFE/REVIEW/BLOCK)"],
        ["Học từ hành vi thực tế", "Chỉ đếm lượt xem", "Không học thói quen nấu", "Không có trạng thái lưu", "Học từ: Pick, Cook, Repeat, Swap, Waste"],
        ["Nấu ăn rảnh tay", "Text từng bước", "Không có", "Phải bấm điện thoại", "Voice-assisted Guided Cooking"]
    ]
    create_table(doc, gap_headers, gap_data, col_widths=[1.5, 1.2, 1.2, 1.2, 1.4])

    # ----------------------------------------------------
    # CHƯƠNG 3: KIẾN TRÚC HỆ THỐNG VÀ NGUYÊN TẮC
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 3: KIẾN TRÚC HỆ THỐNG VÀ NGUYÊN TẮC CÔNG NGHỆ", level=1)
    doc.add_paragraph(
        "Hệ thống EatVibing vận hành trên kiến trúc 3 tầng phân tách độc lập (Presentation Web/PWA & Mobile Roadmap, Business Logic với AI Coach + Trust Layer, và Persistence Supabase PostgreSQL)."
    )

    add_styled_heading(doc, "3.1. Phân định rõ ràng: AI Coach vs. Deterministic Tool", level=2)
    arch_headers = ["Thành phần", "Công nghệ đảm nhiệm", "Phạm vi trách nhiệm trong hệ thống"]
    arch_data = [
        ["EatVibing AI Coach", "LLM Engine + System Prompt + Voice NLP", "Hội thoại tự nhiên Text/Voice, trích xuất nguyên liệu, tạo ứng viên món ăn, tạo lời giải thích minh bạch 'Why recommended?' và gợi ý thay thế nguyên liệu."],
        ["Deterministic Services (Tools)", "Backend Node.js + Supabase SQL", "Tính toán BMI ngữ cảnh, tính Calo/Macro theo trọng lượng, quy đổi đơn vị đo lường, khấu trừ Smart Pantry tự động, và thực thi bộ lọc dị ứng SAFE/REVIEW/BLOCK."]
    ]
    create_table(doc, arch_headers, arch_data, col_widths=[1.8, 2.0, 2.7])

    # ----------------------------------------------------
    # CHƯƠNG 4: MÔ TẢ CHI TIẾT CÁC PHÂN HỆ CHỨC NĂNG
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 4: MÔ TẢ CHI TIẾT CÁC PHÂN HỆ CHỨC NĂNG", level=1)
    
    modules = [
        ("4.1. Hồ sơ Cá nhân hóa Onboarding (Lightweight Profile)",
         "Thu thập thông tin thể trạng cơ bản (Tuổi, Chiều cao, Cân nặng, BMI ngữ cảnh, Mức độ vận động), Mục tiêu ăn uống (Cân bằng, Quản lý cân nặng, Tăng đạm, Tiết kiệm, Giảm lãng phí), Hồ sơ khẩu vị (Cuisine yêu thích, Độ cay, Thói quen ăn uống) và Ràng buộc cứng (Dị ứng, Kiêng khem, Thiết bị bếp, Thời gian nấu, Khẩu phần). Người dùng có toàn quyền sửa, xóa, xuất dữ liệu và Opt-out learning."),
        
        ("4.2. Trợ lý AI Ăn uống Cá nhân hóa (EatVibing AI Coach)",
         "Đóng vai trò chuyên gia tư vấn dinh dưỡng và đầu bếp riêng xuyên suốt hành trình ăn uống. Hỗ trợ hội thoại đa phương thức Text & Voice, tự động giải thích lý do đề xuất ('Why this meal?'), thấu hiểu nguyên liệu tủ lạnh và liên tục học từ phản hồi thực tế của người dùng."),
        
        ("4.3. Không gian Gợi ý Thực đơn Thích ứng (Personalized Meal Guideline)",
         "Không phải thực đơn ép buộc cứng nhắc, mà là không gian gợi ý thích ứng theo ngày/tuần (Sáng, Trưa, Tối, Bữa phụ). Hệ thống học sâu từ toàn bộ chuỗi hành vi: View ➔ Save ➔ Pick ➔ Cook ➔ Skip ➔ Repeat-cook ➔ Rate ➔ Swap ➔ Leftover để tinh chỉnh gợi ý ngày càng chính xác."),
        
        ("4.4. Món Thịnh hành có Kiểm duyệt (Trending / Community Guideline)",
         "Xếp hạng món ăn theo Qualified Engagement Score (Nấu thực tế, nấu lại, đánh giá cao, tỷ lệ bỏ qua thấp). Mọi món trending bắt buộc phải qua kiểm duyệt và đạt chuẩn SAFE. Trending KHÔNG BAO GIỜ được vượt qua các ràng buộc dị ứng (Allergies) của người dùng."),
        
        ("4.5. Lớp Ủy thác Ẩm thực (Culinary Trust Layer)",
         "Cơ chế kiểm soát an toàn mang tính xác định với 3 trạng thái: 🟢 SAFE (Đã xác thực nguyên liệu, dinh dưỡng chuẩn, không dị ứng); 🟡 REVIEW (Thiếu định lượng chi tiết hoặc cần người dùng xác nhận khẩu phần); 🔴 BLOCK (Phát hiện chất gây dị ứng hoặc không an toàn - lập tức chặn khỏi thực đơn)."),
        
        ("4.6. Quản lý Kho Nguyên liệu Thông minh (Smart Pantry)",
         "Theo dõi tồn kho thực tế trong gian bếp bằng Text, Voice hoặc Quét hóa đơn/Barcode (Roadmap). Tự động phát cảnh báo nguyên liệu sắp hết hạn (Near-expiry alerts) và tự động trừ kho nguyên liệu sau khi người dùng nấu xong."),
        
        ("4.7. Hướng dẫn Nấu ăn Rảnh tay Từng bước (Voice-Assisted Guided Cooking)",
         "Hiển thị timeline các bước nấu trực quan kết hợp điều khiển giọng nói rảnh tay ('Next step', 'Set timer 10 minutes', 'Repeat instruction'). Tự động yêu cầu xác nhận lại nếu nhận diện độ tin cậy thấp ở các thông tin nhạy cảm."),
        
        ("4.8. Tạo Danh sách Mua sắm Bù trừ (Smart Shopping List & Diff Engine)",
         "Tự động tính toán lượng nguyên liệu cần mua bằng hiệu số giữa tổng nguyên liệu công thức cần và nguyên liệu sẵn có trong Smart Pantry. Tự động gộp định lượng và phân loại theo quầy siêu thị."),
        
        ("4.9. Hộ chiếu Ẩm thực & Bộ sưu tập (Meal Passport & Recipe Gallery)",
         "Ghi nhận hành trình dinh dưỡng, các món đã nấu thành công, huy hiệu thành tích ẩm thực và lưu lại các biến tấu công thức riêng của người dùng."),
        
        ("4.10. Bảng điều khiển Quản trị viên (Admin Dashboard & CMS)",
         "Cung cấp công cụ quản trị duyệt công thức, kiểm duyệt bài đăng cộng đồng, quản lý kho dữ liệu dinh dưỡng gốc và giám sát các chỉ số an toàn hệ thống.")
    ]
    
    for title, desc in modules:
        add_styled_heading(doc, title, level=2)
        doc.add_paragraph(desc)

    # ----------------------------------------------------
    # CHƯƠNG 5: CƠ SỞ DỮ LIỆU & AN TOÀN
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 5: MÔ HÌNH CƠ SỞ DỮ LIỆU VÀ AN TOÀN RIÊNG TƯ", level=1)
    
    doc.add_paragraph(
        "Hệ thống sử dụng Supabase PostgreSQL với các thực thể cốt lõi: user_profiles (thông tin thể trạng & khẩu vị), smart_pantry (kho nguyên liệu), meals (món ăn & trust_status), ingredients, recipes, interactions (learning log) và messages (lịch sử chat)."
    )
    
    add_callout(doc, "NGUYÊN TẮC PHI Y TẾ HÓA & BẢO VỆ DỮ LIỆU CÁ NHÂN",
        "1. Phi y tế hóa (Non-Medical Disclaimer): EatVibing là công cụ hỗ trợ lối sống ẩm thực; tuyệt đối không đưa ra chẩn đoán y khoa, không thay thế chuyên gia y tế/dinh dưỡng và không dùng BMI đơn lẻ để chẩn đoán.\n"
        "2. Đồng thuận rõ ràng & Thu thập tối thiểu: Mọi dữ liệu chỉ được lưu khi người dùng đồng thuận.\n"
        "3. Quyền làm chủ dữ liệu: Người dùng có quyền sửa, xóa, xuất dữ liệu và bật/tắt chế độ học tập cá nhân hóa (Opt-out Personalization).",
        bg_hex="FEF3C7", border_hex="D97706")

    # ----------------------------------------------------
    # CHƯƠNG 6: CHIẾN LƯỢC KINH DOANH SAAS
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 6: CHIẾN LƯỢC KINH DOANH SẢN PHẨM SỐ (SAAS PRICING)", level=1)
    
    saas_headers = ["Gói dịch vụ", "Giá tham chiếu", "Giá trị cốt lõi cung cấp", "Đối tượng mục tiêu"]
    saas_data = [
        ["Free Tier", "$0", "Smart Pantry cơ bản (20 mục), EatVibing AI Coach cơ bản (Text), Gợi ý thực đơn 3 ngày, Trust Status cơ bản.", "Người dùng mới trải nghiệm core loop."],
        ["Plus", "$4.99 / tháng", "Lập kế hoạch bữa ăn cá nhân hóa với AI Coach, lưu pantry/lịch sử không giới hạn, guideline tuần, cảnh báo gần hết hạn, waste insights và voice cơ bản.", "Cá nhân bận rộn muốn tối ưu hóa đi chợ và nấu ăn."],
        ["Pro", "$9.99 / tháng", "AI Coach cá nhân hóa nâng cao dựa trên profile và hành vi thực tế, phân tích dinh dưỡng sâu, quota voice/vision cao, guideline thích ứng, Meal Passport chi tiết.", "Người theo đuổi chế độ ăn chuyên sâu (Eat Clean, Gym, Keto)."],
        ["Family", "$14.99 / tháng", "Nhiều hồ sơ AI cá nhân hóa trong cùng một hộ gia đình (Household), shared pantry/meal plan/shopping list, phân quyền thành viên và household insights.", "Hộ gia đình có nhiều thành viên khẩu vị khác nhau."],
        ["Creator", "Freemium / RevShare", "Bộ công cụ CMS sáng tạo công thức, xác thực bản quyền (Provenance), chia sẻ doanh thu từ lượt nấu cộng đồng.", "Đầu bếp chuyên nghiệp, Food Blogger."],
        ["Business / API", "Theo hợp đồng B2B", "API tích hợp siêu thị online (WinMart, GrabMart), White-label cho căn hộ/campus, Campaign dashboards, SLA/SSO.", "Doanh nghiệp bán lẻ thực phẩm, Chuỗi phòng gym."]
    ]
    create_table(doc, saas_headers, saas_data, col_widths=[1.2, 1.2, 2.5, 1.6])

    # ----------------------------------------------------
    # CHƯƠNG 7: KỊCH BẢN DEMO 3 PHÚT
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 7: KỊCH BẢN TRÌNH DIỄN SẢN PHẨM (3-MINUTE DEMO SCENARIO)", level=1)
    demo_headers = ["Thời gian", "Phân đoạn trình diễn", "Nội dung & Tác vụ thực hiện", "Thông điệp chứng minh"]
    demo_data = [
        ["0:00 - 0:25", "Onboarding Profile", "Người dùng chọn dị ứng đậu phộng, mục tiêu tăng đạm, thời gian nấu < 30 phút.", "Personalized Profile + Hard Constraints."],
        ["0:25 - 0:45", "Voice Interaction", "Người dùng bấm mic nói: 'Tối nay có trứng, cà chua, cơm, muốn món dưới 30 phút'.", "Voice + Tương tác tự nhiên với AI Coach."],
        ["0:45 - 1:10", "Personalized Guideline", "AI Coach sinh 3 ứng viên món ăn kèm bảng giải thích 'Why recommended?'.", "Personalization + Deterministic Tool."],
        ["1:10 - 1:35", "Culinary Trust Layer", "Món trending dính đậu phộng -> BLOCK; món thiếu data -> REVIEW; món đạt chuẩn -> SAFE.", "Trending KHÔNG ĐƯỢC vượt qua Trust Layer."],
        ["1:35 - 2:00", "Pick món & Shopping List", "User pick món SAFE. Shopping List chỉ hiển thị mua 1 bó hành lá còn thiếu.", "Pantry Auto-diff Engine."],
        ["2:00 - 2:30", "Guided Cooking Voice", "Nấu rảnh tay: ra lệnh 'Next step' và 'Set timer 5 minutes'.", "Hands-free Voice Guided Cooking."],
        ["2:30 - 3:00", "Cooked & Learning Loop", "User đánh dấu Cooked, rating 5 sao: 'Bớt cay'. AI Coach: 'Đã giảm cay cho lần sau'.", "Persistent Personalization Loop."]
    ]
    create_table(doc, demo_headers, demo_data, col_widths=[1.0, 1.5, 2.6, 1.4])

    # ----------------------------------------------------
    # CHƯƠNG 8: LỘ TRÌNH VÀ CHANGELOG
    # ----------------------------------------------------
    add_styled_heading(doc, "CHƯƠNG 8: LỘ TRÌNH PHÁT TRIỂN VÀ CHANGELOG", level=1)
    
    add_styled_heading(doc, "8.1. Lộ trình phát triển (Roadmap)", level=2)
    doc.add_paragraph(
        "- Giai đoạn 1 (Core MVP Prototype): Web/PWA, Profile khảo sát, Smart Pantry cơ bản, EatVibing AI Coach Text, Culinary Trust Layer (SAFE/REVIEW/BLOCK), Shopping List bù trừ.\n"
        "- Giai đoạn 2 (Enhanced Interaction - Q3/2026): Voice NLP hai chiều, Computer Vision quét ảnh tủ lạnh & hóa đơn, Học tập thích ứng dài hạn, Gói Family.\n"
        "- Giai đoạn 3 (Expansion - Q4/2026 - 2027): Native Mobile App (React Native), Creator / Community CMS, B2B Grocery API đặt hàng siêu thị trực tiếp."
    )

    add_styled_heading(doc, "8.2. Changelog Cập nhật Tài liệu (Sau họp nhóm AISC 2026)", level=2)
    cl_headers = ["Hạng mục cập nhật", "Nội dung thay đổi chi tiết"]
    cl_data = [
        ["Terminology & Định vị", "Đổi từ 'AI Culinary Assistant' thành EatVibing AI Coach (Personalized AI Food Coach). Chuẩn hóa vòng lặp Pantry → Plan → Trust → Cook → Learn."],
        ["Onboarding Profile", "Bổ sung khảo sát thể trạng, mục tiêu ăn uống, khẩu vị và ràng buộc cứng (không dùng BMI để chẩn đoán y tế)."],
        ["Personalized Guideline", "Bổ sung không gian gợi ý thích ứng, Explain-Why và học từ hành vi thực tế (Pick, Cook, Repeat, Waste)."],
        ["Trending Guideline", "Xếp hạng bằng Qualified Engagement Score; bắt buộc qua Trust Layer; Trending không được vượt qua Allergen Constraints."],
        ["Voice Interaction Layer", "Bổ sung điều khiển giọng nói khi nấu rảnh tay và nhập pantry kèm cơ chế xác nhận lại khi confidence thấp."],
        ["Định vị Nền tảng", "Phân định rõ Web/PWA là Current Prototype, Native Mobile App là Roadmap Deliverable."],
        ["Mô hình SaaS & Pricing", "Cập nhật 6 gói: Free ($0), Plus ($4.99), Pro ($9.99), Family ($14.99), Creator, Business/API."],
        ["Demo Scenario 3 phút", "Cập nhật kịch bản 3 phút thể hiện trọn vẹn chuỗi giá trị Pantry-to-Plate và Trust Layer."]
    ]
    create_table(doc, cl_headers, cl_data, col_widths=[2.0, 4.5])

    # Save file
    output_path = r"D:\NCKH\EatVibing\THUYET_MINH_DU_AN_EATVIBING.docx"
    doc.save(output_path)
    print(f"File successfully created at: {output_path}")

if __name__ == "__main__":
    build_document()
